"""Prévias do Laboratório: um trecho de cada planilha, o começo dos scripts e o tamanho dos arquivos.

Gera src/previas.json, lido pelo build. Rodar de novo quando o professor trocar um arquivo do laboratório:
    python scripts/preparar-previas.py
Precisa de PyMuPDF (pdf), openpyxl (xlsx) e xlrd (xls).
"""
import json
import os
import urllib.request

import fitz
import openpyxl
import xlrd

BASE = 'https://pantoniofran.github.io/financas-economia-gerencial/'
PASTA = 'arquivos/laboratorio/'
# Aba mais representativa de cada planilha (a primeira nem sempre é: algumas começam com dados do Crystal Ball).
ABA = {
    'planilha-de-indices-financeiros.xlsx': 'Índices',
    'planilha-do-modelo-capm.xls': 'Plan1',
    'planilha-opcoes-reais-em-projeto-solar.xlsx': 'Plan1',
    'planilha-projeto-de-hidrogenio-verde.xlsx': 'Determinístico',
    'planilha-maximizacao-com-tres-produtos.xlsx': 'Planilha1',
}
LINHAS, COLUNAS = 14, 8


def numero(v):
    if isinstance(v, bool) or v is None:
        return '' if v is None else str(v)
    if isinstance(v, (int, float)):
        if float(v).is_integer():
            return f'{int(v):,}'.replace(',', '.')
        casas = 2 if abs(v) >= 1 else 4
        return f'{v:,.{casas}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return str(v).strip()


def recortar(linhas):
    linhas = [[numero(v) for v in l[:COLUNAS]] for l in linhas]
    linhas = [l for l in linhas if any(l)][:LINHAS]
    usadas = [c for c in range(COLUNAS) if any(c < len(l) and l[c] for l in linhas)]
    return [[l[c] if c < len(l) else '' for c in usadas] for l in linhas]


def tamanho(caminho):
    b = os.path.getsize(caminho)
    return f'{b / 1048576:.1f} MB'.replace('.', ',') if b >= 1048576 else f'{max(1, round(b / 1024))} KB'


previas = {}
for nome in sorted(os.listdir(PASTA)):
    caminho = PASTA + nome
    url = BASE + caminho
    p = {'tamanho': tamanho(caminho)}
    if nome.endswith('.xlsx'):
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        aba = ABA.get(nome, wb.sheetnames[0])
        p.update(tipo='tabela', abas=wb.sheetnames, aba=aba,
                 linhas=recortar(list(wb[aba].iter_rows(max_row=60, max_col=COLUNAS, values_only=True))))
    elif nome.endswith('.xls'):
        wb = xlrd.open_workbook(caminho)
        aba = ABA.get(nome, wb.sheet_names()[0])
        s = wb.sheet_by_name(aba)
        p.update(tipo='tabela', abas=wb.sheet_names(), aba=aba,
                 linhas=recortar([s.row_values(r) for r in range(min(s.nrows, 60))]))
    elif nome.endswith('.pdf'):
        p.update(paginas=fitz.open(caminho).page_count)
    elif nome.endswith('.txt'):
        with open(caminho, encoding='utf-8') as f:
            p.update(tipo='codigo', linhas=f.read().splitlines()[:40])
    previas[url] = p

# Repositório do professor no GitHub: as pastas viram a prévia.
REPO = 'https://github.com/pantoniofran/GitSeminarioFinancas.git'
try:
    with urllib.request.urlopen('https://api.github.com/repos/pantoniofran/GitSeminarioFinancas/contents', timeout=20) as r:
        itens = json.load(r)
    previas[REPO] = {'tipo': 'pastas', 'itens': [i['name'] for i in itens if i['type'] == 'dir']}
except OSError as erro:
    print('repositório não lido:', erro)

with open('src/previas.json', 'w', encoding='utf-8') as f:
    json.dump(previas, f, ensure_ascii=False, indent=2)
    f.write('\n')
print(f'{len(previas)} prévias em src/previas.json')
